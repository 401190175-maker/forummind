"""Format parsing and deterministic document chunking contracts."""

from docx import Document
from docx.shared import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.documents.chunker import chunk_pages
from app.documents.parsers import parse_document


def make_pdf(path):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
    })
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 18 Tf 72 720 Td (compressive strength) Tj ET")
    page[NameObject("/Contents")] = stream
    with path.open("wb") as handle:
        writer.write(handle)


def make_docx(path):
    document = Document()
    document.add_paragraph("Foam density and compressive strength")
    document.add_paragraph("The second observation remains traceable.")
    document.save(path)


def test_parse_pdf_keeps_page_location(tmp_path):
    path = tmp_path / "study.pdf"
    make_pdf(path)

    pages = parse_document(path, "application/pdf")

    assert pages[0].page_or_location == "page:1"
    assert "compressive strength" in pages[0].content


def test_parse_docx_text_markdown_and_plain_text(tmp_path):
    docx_path = tmp_path / "study.docx"
    make_docx(docx_path)
    text_path = tmp_path / "study.txt"
    text_path.write_text("line one\nline two", encoding="utf-8")
    markdown_path = tmp_path / "study.md"
    markdown_path.write_text("# Results\n\nThe result is reproducible.", encoding="utf-8")

    docx_pages = parse_document(docx_path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    text_pages = parse_document(text_path, "text/plain")
    markdown_pages = parse_document(markdown_path, "text/markdown")

    assert docx_pages[0].page_or_location == "paragraph:1"
    assert "Foam density" in " ".join(page.content for page in docx_pages)
    assert text_pages[0].page_or_location == "document:1"
    assert markdown_pages[0].page_or_location == "document:1"


def test_parse_rejects_empty_text_file(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("\n", encoding="utf-8")

    from app.documents.parsers import DocumentParseError

    try:
        parse_document(path, "text/plain")
    except DocumentParseError as exc:
        assert "empty" in str(exc)
    else:
        raise AssertionError("empty document must be rejected")


def test_chunking_is_stable_and_keeps_location_metadata():
    parsed = [
        type("Page", (), {
            "page_or_location": "page:1",
            "content": "0123456789" * 4,
        })(),
    ]

    first = chunk_pages(parsed, max_chars=12, overlap=3, document_id="doc-1")
    second = chunk_pages(parsed, max_chars=12, overlap=3, document_id="doc-1")

    assert [chunk.chunk_id for chunk in first] == [chunk.chunk_id for chunk in second]
    assert first[0].page_or_location == "page:1"
    assert first[0].char_start == 0
    assert first[0].char_end == 12
    assert all(0 < len(chunk.content) <= 12 for chunk in first)
