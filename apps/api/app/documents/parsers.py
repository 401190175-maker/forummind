"""Format-specific extraction for uploaded research documents."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class DocumentParseError(ValueError):
    """The document cannot be parsed into ordered text pages."""


@dataclass(frozen=True)
class ParsedPage:
    page_or_location: str
    content: str


PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT_MIMES = {"text/plain", "text/markdown", "text/x-markdown"}


def parse_document(path: str | Path, mime_type: str) -> list[ParsedPage]:
    """Extract ordered, non-empty text blocks from one supported document."""
    source = Path(path)
    mime = mime_type.strip().lower()
    try:
        if mime == PDF_MIME:
            pages = _parse_pdf(source)
        elif mime == DOCX_MIME:
            pages = _parse_docx(source)
        elif mime in TEXT_MIMES:
            pages = _parse_text(source)
        else:
            raise DocumentParseError(f"unsupported document MIME type: {mime_type}")
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"document could not be parsed: {exc}") from exc
    if not pages or not any(page.content.strip() for page in pages):
        raise DocumentParseError("document is empty")
    return pages


def _parse_pdf(path: Path) -> list[ParsedPage]:
    reader = PdfReader(str(path))
    pages: list[ParsedPage] = []
    for index, page in enumerate(reader.pages, start=1):
        content = (page.extract_text() or "").strip()
        if content:
            pages.append(ParsedPage(page_or_location=f"page:{index}", content=content))
    return pages


def _parse_docx(path: Path) -> list[ParsedPage]:
    document = Document(str(path))
    pages: list[ParsedPage] = []
    location = 0
    for paragraph in document.paragraphs:
        content = paragraph.text.strip()
        if content:
            location += 1
            pages.append(ParsedPage(page_or_location=f"paragraph:{location}", content=content))
    for table_index, table in enumerate(document.tables, start=1):
        content = "\n".join(
            "\t".join(cell.text.strip() for cell in row.cells)
            for row in table.rows
        ).strip()
        if content:
            location += 1
            pages.append(ParsedPage(page_or_location=f"table:{table_index}", content=content))
    return pages


def _parse_text(path: Path) -> list[ParsedPage]:
    content = path.read_text(encoding="utf-8").replace("\r\n", "\n").strip()
    return [ParsedPage(page_or_location="document:1", content=content)] if content else []
